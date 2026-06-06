"""Introspect + dump thruuu .xlsx exports for the dropshipping audit."""
import csv
import os
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SRC = r"D:\TÉLÉCHARGEMENT"
FILES = [
    "thruuu_export_dropshipping shopify vs wordpress_2026-5-29.xlsx",
    "thruuu_export_shopify vs woocommerce_2026-5-29.xlsx",
    "thruuu_export_dropshipping woocommerce_2026-5-29.xlsx",
]
OUT = Path(__file__).parent

def cell(v):
    if v is None:
        return ""
    return str(v).replace("\r", " ").replace("\n", " ").strip()

for fname in FILES:
    path = os.path.join(SRC, fname)
    if not os.path.exists(path):
        print(f"!! MISSING: {path}")
        continue
    slug = fname.replace("thruuu_export_", "").replace("_2026-5-29.xlsx", "").replace(" ", "-")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f"\n========== {fname} ==========")
    print("SHEETS:", wb.sheetnames)
    for ws in wb.worksheets:
        print(f"\n--- sheet '{ws.title}' : {ws.max_row} rows x {ws.max_column} cols ---")
        # dump full sheet to csv
        dump = OUT / f"{slug}__{ws.title.replace('/', '-').replace(' ', '-')}.csv"
        rows = list(ws.iter_rows(values_only=True))
        with open(dump, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for r in rows:
                w.writerow([cell(c) for c in r])
    wb.close()
print("\nDONE")
