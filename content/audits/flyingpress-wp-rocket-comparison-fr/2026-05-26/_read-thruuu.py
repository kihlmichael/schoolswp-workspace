"""Extract thruuu audit (.docx) + SERP analysis (.xlsx) into JSON for both languages."""
import json
from pathlib import Path
import docx
import openpyxl

DL = Path(r"D:\TÉLÉCHARGEMENT")
OUT = Path(__file__).resolve().parent

DOCX_FILES = {
    "audit_fr": DL / "Expor_audit_6a1567086d4d2c2ae024950b.docx",
    "audit_de": DL / "Expor_audit_6a1566e96d4d2c2ae02494c5.docx",
}

XLSX_FILES = {
    "serp_main": DL / "thruuu_export_flyingpress vs wp rocket_2026-5-26.xlsx",
    "serp_dup": DL / "thruuu_export_flyingpress vs wp rocket_2026-5-26 (1).xlsx",
}


def read_docx(path: Path) -> dict:
    d = docx.Document(str(path))
    paragraphs = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = p.style.name if p.style else "Normal"
        paragraphs.append({"style": style, "text": text})
    tables = []
    for t in d.tables:
        rows = []
        for row in t.rows:
            rows.append([c.text.strip() for c in row.cells])
        tables.append(rows)
    return {"paragraphs": paragraphs, "tables": tables, "n_paragraphs": len(paragraphs), "n_tables": len(tables)}


def read_xlsx(path: Path) -> dict:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True, max_row=200):
            if any(c is not None for c in row):
                rows.append([str(c) if c is not None else "" for c in row])
        out[sheet_name] = {
            "n_rows": len(rows),
            "n_cols_first_row": len(rows[0]) if rows else 0,
            "rows": rows,
        }
    wb.close()
    return out


payload = {"docx": {}, "xlsx": {}}
for key, path in DOCX_FILES.items():
    try:
        payload["docx"][key] = read_docx(path)
    except Exception as e:
        payload["docx"][key] = {"error": str(e)}

for key, path in XLSX_FILES.items():
    try:
        payload["xlsx"][key] = read_xlsx(path)
    except Exception as e:
        payload["xlsx"][key] = {"error": str(e)}

out_file = OUT / "_thruuu-raw.json"
out_file.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"WROTE {out_file} ({out_file.stat().st_size} bytes)")
print("DOCX keys:", list(payload["docx"].keys()))
for k, v in payload["docx"].items():
    if isinstance(v, dict) and "n_paragraphs" in v:
        print(f"  {k}: {v['n_paragraphs']} paragraphs, {v['n_tables']} tables")
print("XLSX keys:", list(payload["xlsx"].keys()))
for k, v in payload["xlsx"].items():
    if isinstance(v, dict) and "error" not in v:
        sheets_summary = ", ".join(f"{s}({d['n_rows']}r)" for s, d in v.items())
        print(f"  {k}: sheets = {sheets_summary}")
