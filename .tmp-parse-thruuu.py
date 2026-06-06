"""Parse the thruuu xlsx export + list candidate draft files in TELECHARGEMENT."""
import glob
import os

DL = "D:/TÉLÉCHARGEMENT"

print("=== Files in TELECHARGEMENT matching fluentcart/woocommerce/thruuu ===")
for f in glob.glob(DL + "/*"):
    name = os.path.basename(f)
    low = name.lower()
    if any(k in low for k in ("fluentcart", "woocommerce", "thruuu", "vs woo")):
        size = os.path.getsize(f)
        print(f"  {name}  ({size/1024:.0f} KB)")

xlsx = None
for f in glob.glob(DL + "/*.xlsx"):
    if "fluentcart" in os.path.basename(f).lower():
        xlsx = f
        break

print("\n=== thruuu xlsx:", xlsx, "===")
if xlsx:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        print("SHEETS:", wb.sheetnames)
        for ws in wb.worksheets:
            print(f"\n--- SHEET: {ws.title}  (dims {ws.max_row} x {ws.max_column}) ---")
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 12:
                    break
                cells = [str(c)[:60] if c is not None else "" for c in row]
                rows.append(" | ".join(cells))
            for r in rows:
                print("  " + r[:300])
    except ImportError as e:
        print("openpyxl missing:", e)
